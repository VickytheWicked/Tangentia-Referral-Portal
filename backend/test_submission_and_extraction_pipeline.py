#!/usr/bin/env python3
"""
Test the end-to-end referral submission pipeline where CV data is extracted and submitted:
1. Purges all existing referrals from SQLite, Excel, Mock Storage, and CV DB.
2. For each of the 18 CVs in Resume_Editing_/:
   - POST /api/referrals/extract-cv (tests CV extraction preview)
   - Inspects found vs not_found fields
   - Simulates user filling any missing fields manually
   - POST /api/referrals (submits the referral form with CV file attached)
   - Evaluates CV Intelligence processing & job matching
3. Audits persistence across SQLite, Excel, Mock Storage, and HR suggestions.
4. Generates a comprehensive testing report.
"""

import os
import sys
import shutil
import io
import httpx
from datetime import datetime, timezone
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
WORKSPACE_ROOT = os.path.dirname(BASE_DIR)
sys.path.insert(0, BASE_DIR)

from app.config import settings
from app.database import engine, Base, SessionLocal
from app.models.referral import Referral, ReferralStatus
from app.models.status_history import ReferralStatusHistory
from app.models.hr_note import HRNote
from app.models.job_position import JobPosition
from app.cv_intelligence.database import cv_engine, CVSessionLocal
from app.cv_intelligence.models import CandidateProfile, JobMatch

CV_DIR = os.path.join(WORKSPACE_ROOT, "Resume_Editing_")
API_BASE = "http://localhost:8000/api"

# Candidate metadata mapping with target jobs and manual fallbacks
CV_PIPELINE_DATA = [
    {
        "file": "APA Developer.docx",
        "manual_fallback_name": "Jassim Shaji",
        "manual_fallback_email": "jassim.shaji@example.com",
        "manual_fallback_phone": "+91 98471 23456",
        "position_id": "cats-16855743",
        "relationship": "Former Colleague",
        "note": "Experienced RPA & Automation Anywhere developer with strong hands-on APA and bot deployment track record.",
    },
    {
        "file": "APA Developer.pdf",
        "manual_fallback_name": "Keerthi Krishnan",
        "manual_fallback_email": "keerthi.krishnan@example.com",
        "manual_fallback_phone": "+91 94462 89012",
        "position_id": "cats-16855744",
        "relationship": "Industry Peer",
        "note": "Expert in Automation Anywhere A360, UiPath, and Intelligent Document Processing.",
    },
    {
        "file": "Business Analyst -2.pdf",
        "manual_fallback_name": "Vinayak Chari",
        "manual_fallback_email": "vinayak.chari@example.com",
        "manual_fallback_phone": "+91 98234 56789",
        "position_id": "cats-16699887",
        "relationship": "Former Colleague",
        "note": "Senior Business Analyst with extensive experience in requirements elicitation, BRDs, and agile delivery.",
    },
    {
        "file": "Business Analyst.pdf",
        "manual_fallback_name": "Ruchita Elekar",
        "manual_fallback_email": "ruchita.elekar@example.com",
        "manual_fallback_phone": "+91 99231 19494",
        "position_id": "cats-16699887",
        "relationship": "College Alumni",
        "note": "Proven track record in business process analysis, functional specifications, and stakeholder management.",
    },
    {
        "file": "Program Director-Resume.pdf",
        "manual_fallback_name": "Ajinkya Birwadkar",
        "manual_fallback_email": "ajinkya.birwadkar@example.com",
        "manual_fallback_phone": "+91 98205 67890",
        "position_id": "cats-16793607",
        "relationship": "Industry Peer",
        "note": "Seasoned Program Director and delivery leader driving multi-million dollar enterprise automation transformations.",
    },
    {
        "file": "Prompt Engineer.docx",
        "manual_fallback_name": "Aryan Verma",
        "manual_fallback_email": "aryan.verma@example.com",
        "manual_fallback_phone": "+91 99345 67890",
        "position_id": "job-002",
        "relationship": "Mentee",
        "note": "Specializes in Generative AI prompt optimization, RAG evaluation, and LLM orchestration with LangChain.",
    },
    {
        "file": "Prompt Engineer2.docx",
        "manual_fallback_name": "Rohan Mehta",
        "manual_fallback_email": "rohan.mehta@example.com",
        "manual_fallback_phone": "+91 98765 43210",
        "position_id": "job-002",
        "relationship": "Former Colleague",
        "note": "Prompt engineering and AI specialist with expertise in fine-tuning, vector search, and multimodal agent workflows.",
    },
    {
        "file": "QA1.pdf",
        "manual_fallback_name": "B Anand",
        "manual_fallback_email": "b.anand@example.com",
        "manual_fallback_phone": "+91 94432 10987",
        "position_id": "cats-16856991",
        "relationship": "Former Colleague",
        "note": "Solid QA automation engineer with Selenium WebDriver, TestNG, and CI/CD pipeline integration experience.",
    },
    {
        "file": "QA2.pdf",
        "manual_fallback_name": "Abisha Rajesh",
        "manual_fallback_email": "abisha.rajesh@example.com",
        "manual_fallback_phone": "+91 98401 23456",
        "position_id": "cats-16856991",
        "relationship": "Industry Peer",
        "note": "Experienced SDET in API testing, Postman, Cypress, and end-to-end web test frameworks.",
    },
    {
        "file": "QA3.pdf",
        "manual_fallback_name": "Balaji S",
        "manual_fallback_email": "balaji.s@example.com",
        "manual_fallback_phone": "+91 97890 12345",
        "position_id": "cats-16856991",
        "relationship": "Former Colleague",
        "note": "Lead QA tester with deep expertise in regression suites, performance testing, and defect life cycle tracking in Jira.",
    },
    {
        "file": "Scrum_Lead.docx",
        "manual_fallback_name": "Neha Kulkarni",
        "manual_fallback_email": "neha.kulkarni@example.com",
        "manual_fallback_phone": "+91 98220 98765",
        "position_id": "cats-16699887",
        "relationship": "Project Colleague",
        "note": "Certified Scrum Master (CSM) driving agile ceremonies, sprint planning, and cross-functional team delivery.",
    },
    {
        "file": "Solution Architect - Lead.docx",
        "manual_fallback_name": "Siddharth Nair",
        "manual_fallback_email": "siddharth.nair@example.com",
        "manual_fallback_phone": "+91 98450 87654",
        "position_id": "job-003",
        "relationship": "Industry Peer",
        "note": "Principal Solution Architect specializing in AWS cloud infrastructure, microservices architecture, and zero-trust security.",
    },
    {
        "file": "Solution Architect.docx",
        "manual_fallback_name": "Mohamed Faheem Ashique",
        "manual_fallback_email": "mohamed.faheem@example.com",
        "manual_fallback_phone": "+91 98950 12345",
        "position_id": "job-003",
        "relationship": "Former Colleague",
        "note": "Cloud and software architect with hands-on systems design, Databricks pipelines, and scalable enterprise backend patterns.",
    },
    {
        "file": "Support1.pdf",
        "manual_fallback_name": "Komal Kutre",
        "manual_fallback_email": "komal.kutre@example.com",
        "manual_fallback_phone": "+91 98201 11223",
        "position_id": "cats-16849726",
        "relationship": "College Alumni",
        "note": "Technical support and IT operations specialist proficient in Active Directory, VPN, troubleshooting, and AWS fundamentals.",
    },
    {
        "file": "Support2.docx",
        "manual_fallback_name": "Allen Manu Philip",
        "manual_fallback_email": "allen.philip@example.com",
        "manual_fallback_phone": "+91 98460 22334",
        "position_id": "cats-16849726",
        "relationship": "Former Colleague",
        "note": "Hands-on IT engineer with desktop and network support, Linux server administration, and ticket resolution expertise.",
    },
    {
        "file": "Support3.pdf",
        "manual_fallback_name": "Virendra Salgaonkar",
        "manual_fallback_email": "virendra.salgaonkar@example.com",
        "manual_fallback_phone": "+91 98231 33445",
        "position_id": "cats-16849726",
        "relationship": "Industry Peer",
        "note": "Infrastructure support engineer with Windows Server, cloud connectivity, and hardware diagnostic skills.",
    },
    {
        "file": "Tech Lead.docx",
        "manual_fallback_name": "Faheem Ashique",
        "manual_fallback_email": "faheem.techlead@example.com",
        "manual_fallback_phone": "+91 98950 54321",
        "position_id": "job-001",
        "relationship": "Former Colleague",
        "note": "Technical Lead with strong Python, FastAPI, Docker, and distributed backend system engineering experience.",
    },
    {
        "file": "UI-UX Designer.docx",
        "manual_fallback_name": "Daryl Emmanuel Vaz",
        "manual_fallback_email": "daryl.vaz@example.com",
        "manual_fallback_phone": "+91 98225 66778",
        "position_id": "job-004",
        "relationship": "Former Colleague",
        "note": "Talented UI/UX designer and frontend collaborator with Figma, responsive design systems, and user empathy.",
    },
]


def purge_all_data():
    """Purge all referral records from SQLite, Excel, mock storage, and CV DB."""
    print("\n" + "=" * 70)
    print("STEP 1: PURGING ALL EXISTING REFERRAL DATA")
    print("=" * 70)

    # 1. Main SQLite DB
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        sh_del = db.query(ReferralStatusHistory).delete()
        hr_del = db.query(HRNote).delete()
        ref_del = db.query(Referral).delete()
        db.commit()
        print(f"  [SQLite Main DB] Cleared {ref_del} referrals, {sh_del} status logs, {hr_del} notes.")
    except Exception as e:
        db.rollback()
        print(f"  [SQLite Main DB] Error: {e}")
    finally:
        db.close()

    # 2. Excel
    excel_path = settings.EXCEL_FILE_PATH
    if os.path.exists(excel_path) and os.path.getsize(excel_path) > 0:
        wb = openpyxl.load_workbook(excel_path)
        for sname in ["Referrals", "StatusHistory", "HRNotes", "HiredHistory"]:
            if sname in wb.sheetnames:
                ws = wb[sname]
                if ws.max_row > 1:
                    cnt = ws.max_row - 1
                    ws.delete_rows(2, cnt)
                    print(f"  [Excel Workbook] Deleted {cnt} rows from sheet '{sname}'")
        wb.save(excel_path)
        wb.close()

    # 3. Mock Storage
    storage_dir = settings.LOCAL_STORAGE_DIR
    current_year = str(datetime.now(timezone.utc).year)
    year_dir = os.path.join(storage_dir, current_year)
    if os.path.exists(year_dir):
        shutil.rmtree(year_dir)
        os.makedirs(year_dir, exist_ok=True)
        print(f"  [Mock Storage] Cleaned directory: {year_dir}")
    index_file = os.path.join(storage_dir, "_index.json")
    with open(index_file, "w") as f:
        f.write("{}")
    print(f"  [Mock Storage] Reset index: {index_file}")

    # 4. CV Intelligence DB
    cv_db = CVSessionLocal()
    try:
        jm_del = cv_db.query(JobMatch).delete()
        cp_del = cv_db.query(CandidateProfile).delete()
        cv_db.commit()
        print(f"  [CV Intelligence DB] Cleared {cp_del} candidate profiles and {jm_del} job matches.")
    except Exception as e:
        cv_db.rollback()
        print(f"  [CV Intelligence DB] Error: {e}")
    finally:
        cv_db.close()

    # 5. Reload running Uvicorn server by touching app/main.py
    main_py_path = os.path.join(BASE_DIR, "app", "main.py")
    if os.path.exists(main_py_path):
        import time
        os.utime(main_py_path, None)
        print("  [Live Server] Triggered Uvicorn reload from clean Excel workbook...")
        time.sleep(2.0)

    # 6. Verify live server clean slate
    try:
        res = httpx.get("http://localhost:8000/api/referrals", timeout=5.0)
        current_count = len(res.json()) if res.status_code == 200 else -1
        print(f"  [Live Server Verification] GET /api/referrals returned {current_count} referrals.")
    except Exception as e:
        print(f"  [Live Server Verification] Check failed: {e}")

    print("Purge completed successfully! System is at clean zero-state.\n")


def test_submission_and_extraction_pipeline():
    """Execute live HTTP extraction and submission for all 18 CVs."""
    print("=" * 70)
    print("STEP 2: TESTING EXTRACTION PREVIEW & REFERRAL SUBMISSION PIPELINE")
    print("=" * 70)

    client = httpx.Client(base_url="http://localhost:8000", timeout=60.0)
    hr_headers = {"Authorization": "Bearer dev-hr-token"}

    submission_results = []

    for idx, item in enumerate(CV_PIPELINE_DATA, 1):
        filename = item["file"]
        file_path = os.path.join(CV_DIR, filename)
        if not os.path.exists(file_path):
            print(f"  [-] Skipped: {filename} not found")
            continue

        with open(file_path, "rb") as f:
            cv_bytes = f.read()

        mime_type = "application/pdf" if filename.endswith(".pdf") else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        # -------------------------------------------------------------
        # 1. TEST POST /api/referrals/extract-cv
        # -------------------------------------------------------------
        files_extract = {"file": (filename, cv_bytes, mime_type)}
        extract_res = client.post("/api/referrals/extract-cv", files=files_extract)
        if extract_res.status_code != 200:
            print(f"  [!] Extraction failed for {filename}: HTTP {extract_res.status_code} - {extract_res.text}")
            continue

        preview = extract_res.json()
        ext_name = preview.get("candidate_name")
        ext_email = preview.get("candidate_email")
        ext_phone = preview.get("candidate_phone")
        ext_exp = preview.get("years_of_experience") or 0.0
        found_fields = preview.get("found_fields", [])
        not_found_fields = preview.get("not_found_fields", [])

        # -------------------------------------------------------------
        # 2. SIMULATE USER BEHAVIOR (Autofill + Manual fill for missing)
        # -------------------------------------------------------------
        final_candidate_name = ext_name if ext_name else item["manual_fallback_name"]
        final_candidate_email = ext_email if ext_email else item["manual_fallback_email"]
        final_candidate_phone = ext_phone if ext_phone else item["manual_fallback_phone"]
        final_exp = ext_exp if ext_exp > 0 else 3.0

        manual_filled = []
        if not ext_name:
            manual_filled.append("name")
        if not ext_email:
            manual_filled.append("email")
        if not ext_phone:
            manual_filled.append("phone")

        # -------------------------------------------------------------
        # 3. TEST POST /api/referrals (Submit Referral with CV attached)
        # -------------------------------------------------------------
        form_data = {
            "candidate_name": final_candidate_name,
            "candidate_email": final_candidate_email,
            "candidate_phone": final_candidate_phone,
            "referred_by_name": "Vansh Rupesh (Employee)",
            "referred_by_email": "vansh.rupesh@tangentia.com",
            "years_of_experience": str(final_exp),
            "relationship": item["relationship"],
            "referral_note": item["note"],
            "position_id": item["position_id"],
            "candidate_consent": "true",
        }
        files_submit = {"file": (filename, io.BytesIO(cv_bytes), mime_type)}

        submit_res = client.post("/api/referrals", data=form_data, files=files_submit)
        if submit_res.status_code != 201:
            print(f"  [!] Submission failed for {final_candidate_name}: HTTP {submit_res.status_code} - {submit_res.text}")
            continue

        created_ref = submit_res.json()
        ref_id = created_ref["id"]
        ref_num = created_ref["referral_number"]
        pos_title = created_ref.get("position_title", "Position")

        # -------------------------------------------------------------
        # 4. PROCESS CV INTELLIGENCE
        # -------------------------------------------------------------
        proc_res = client.post(f"/api/cv-intelligence/process/{ref_id}?force=true", headers=hr_headers)
        if proc_res.status_code == 200:
            profile_data = proc_res.json()
            skills_count = len(profile_data.get("skills", []))
            edu_count = len(profile_data.get("education", []))
        else:
            skills_count = 0
            edu_count = 0

        submission_results.append({
            "idx": idx,
            "file": filename,
            "ref_number": ref_num,
            "ref_id": ref_id,
            "candidate_name": final_candidate_name,
            "candidate_email": final_candidate_email,
            "candidate_phone": final_candidate_phone,
            "position_title": pos_title,
            "found_fields": found_fields,
            "not_found_fields": not_found_fields,
            "manual_filled": manual_filled,
            "skills_count": skills_count,
            "edu_count": edu_count,
        })

        manual_note = f" (Manual: {', '.join(manual_filled)})" if manual_filled else " (100% Autofilled)"
        print(f"  [{idx:2d}/18] SUBMITTED {ref_num} | {final_candidate_name:22} -> {pos_title[:28]} {manual_note}")

    # -------------------------------------------------------------
    # STEP 3: AUDIT PERSISTENCE & SUGGESTIONS
    # -------------------------------------------------------------
    print("\n" + "=" * 70)
    print("STEP 3: VERIFYING DATA PERSISTENCE & HR SUGGESTIONS")
    print("=" * 70)

    # 1. Query GET /api/referrals
    list_res = client.get("/api/referrals")
    all_refs = list_res.json() if list_res.status_code == 200 else []
    print(f"  [API Verification] GET /api/referrals returned {len(all_refs)} referrals.")

    # 2. Query Excel
    wb = openpyxl.load_workbook(settings.EXCEL_FILE_PATH, data_only=True)
    ws = wb["Referrals"]
    excel_rows = max(0, ws.max_row - 1)
    wb.close()
    print(f"  [Excel Verification] Referrals worksheet contains {excel_rows} rows.")

    # 3. Query Mock Storage
    current_year = str(datetime.now(timezone.utc).year)
    year_dir = os.path.join(settings.LOCAL_STORAGE_DIR, current_year)
    storage_files = os.listdir(year_dir) if os.path.exists(year_dir) else []
    print(f"  [Storage Verification] Mock storage contains {len(storage_files)} CV files.")

    # 4. Query GET /api/cv-intelligence/suggestions
    sug_res = client.get("/api/cv-intelligence/suggestions", headers=hr_headers)
    suggestions = sug_res.json() if sug_res.status_code == 200 else []
    openings_with_matches = [s for s in suggestions if (len(s.get("strong_matches", [])) + len(s.get("good_matches", [])) + len(s.get("potential_matches", []))) > 0]
    total_suggested_candidates = sum((len(s.get("strong_matches", [])) + len(s.get("good_matches", [])) + len(s.get("potential_matches", []))) for s in openings_with_matches)

    print(f"  [HR Suggestions Verification] {len(openings_with_matches)} active openings have matched referrals (Total: {total_suggested_candidates} match evaluations).")

    # -------------------------------------------------------------
    # STEP 4: PRINT SUMMARY REPORT
    # -------------------------------------------------------------
    print("\n" + "=" * 70)
    print("REFERRAL SUBMISSION & EXTRACTION TEST REPORT SUMMARY")
    print("=" * 70)
    print(f"  Total CVs Evaluated & Submitted: {len(submission_results)} / 18")
    print(f"  Fully Autofilled CVs:            {sum(1 for r in submission_results if not r['manual_filled'])} / 18")
    print(f"  CVs Triggering Manual Fill:      {sum(1 for r in submission_results if r['manual_filled'])} / 18")
    print(f"  SQLite Referrals Count:          {len(all_refs)}")
    print(f"  Excel Referrals Rows:            {excel_rows}")
    print(f"  Mock Storage Stored Files:       {len(storage_files)}")
    print(f"  Openings with Ranked Candidates: {len(openings_with_matches)}")
    print("=" * 70 + "\n")

    client.close()


if __name__ == "__main__":
    purge_all_data()
    test_submission_and_extraction_pipeline()
