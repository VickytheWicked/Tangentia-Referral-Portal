import os
import io
import logging
import threading
from datetime import datetime, timezone
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import settings
from app.services.excel.base import ExcelServiceInterface

logger = logging.getLogger(__name__)

# Standard Styling Constants
HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Segoe UI", size=10)
BORDER_THIN = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

SHEET_SCHEMAS = {
    "Referrals": [

        "Referral Number",
        "Candidate Name",
        "Candidate Email",
        "Candidate Phone",
        "Referred By",
        "Years Experience",
        "Relationship",
        "Position Title",
        "Status",
        "Referral ID",
        "Position ID",
        "Referred By User ID",
        "LinkedIn URL",
        "GitHub URL",
        "Original CV Filename",
        "SharePoint CV URL",
        "Referral Note",
        "Created At",
        "Updated At",
    ],
    "JobPositions": [
        "Job ID",
        "Title",
        "Department",
        "Location",
        "Employment Type",
        "Is Active",
        "Description",
        "Created At",
        "Updated At",
    ],
    "StatusHistory": [
        "History ID",
        "Referral ID",
        "Referral Number",
        "Old Status",
        "New Status",
        "Changed By",
        "Comment",
        "Created At",
    ],
    "HRNotes": [
        "Note ID",
        "Referral ID",
        "Referral Number",
        "Created By",
        "Note",
        "Created At",
        "Updated At",
    ],
    "Users": [
        "User ID",
        "Name",
        "Email",
        "Password",
        "Role",
        "Department",
        "Created At",
    ],
    "HiredHistory": [
        "Referral Number",
        "Candidate Name",
        "Candidate Email",
        "Position Title",
        "Department",
        "Location",
        "Employment Type",
        "Referred By",
        "Hired Date",
        "Status",
        "Referral ID",
        "Position ID",
    ],
}



class LocalExcelService(ExcelServiceInterface):
    """
    Local Microsoft Excel (.xlsx) implementation using openpyxl.
    Manages Tangentia_Referrals.xlsx with styled sheets, auto-widths, and atomic updates.
    """

    def __init__(self, file_path: Optional[str] = None):
        self.file_path = file_path or settings.EXCEL_FILE_PATH
        self._lock = threading.Lock()
        os.makedirs(os.path.dirname(self.file_path), exist_ok=True)

    def _style_header_row(self, ws, headers: List[str]):
        """Apply Tangentia branding to header row with frozen panes"""
        ws.append(headers)
        ws.freeze_panes = "A2"
        for col_idx, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_THIN
        ws.row_dimensions[1].height = 26

    def _auto_adjust_column_widths(self, ws):
        """Auto-adjust column widths for optimal reading in Excel"""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    def initialize_workbook(self) -> None:
        """Create structured Excel workbook if it doesn't exist, and ensure sheet schemas are current"""
        with self._lock:
            if os.path.exists(self.file_path):
                # Ensure Users sheet exists and has updated headers including Password
                try:
                    wb = openpyxl.load_workbook(self.file_path)
                    modified = False
                    if "Users" not in wb.sheetnames:
                        ws = wb.create_sheet(title="Users")
                        self._style_header_row(ws, SHEET_SCHEMAS["Users"])
                        self._auto_adjust_column_widths(ws)
                        modified = True
                    else:
                        ws = wb["Users"]
                        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
                        if "Password" not in headers:
                            # Re-create/update header row
                            ws.delete_rows(1)
                            ws.insert_rows(1)
                            for col_idx, h in enumerate(SHEET_SCHEMAS["Users"], 1):
                                cell = ws.cell(row=1, column=col_idx, value=h)
                                cell.fill = HEADER_FILL
                                cell.font = HEADER_FONT
                                cell.alignment = Alignment(horizontal="center", vertical="center")
                                cell.border = BORDER_THIN
                            ws.row_dimensions[1].height = 26
                            self._auto_adjust_column_widths(ws)
                            modified = True
                    # Ensure HiredHistory sheet exists
                    if "HiredHistory" not in wb.sheetnames:
                        ws_hh = wb.create_sheet(title="HiredHistory")
                        self._style_header_row(ws_hh, SHEET_SCHEMAS["HiredHistory"])
                        self._auto_adjust_column_widths(ws_hh)
                        modified = True

                    if modified:
                        wb.save(self.file_path)
                    wb.close()
                except Exception as e:
                    logger.warning(f"Error checking sheets in Excel workbook: {e}")
                return


            wb = openpyxl.Workbook()
            default_sheet = wb.active
            wb.remove(default_sheet)

            for sheet_name, headers in SHEET_SCHEMAS.items():
                ws = wb.create_sheet(title=sheet_name)
                self._style_header_row(ws, headers)
                self._auto_adjust_column_widths(ws)

            wb.save(self.file_path)
            logger.info(f"Initialized Microsoft Excel workbook at: {self.file_path}")

    def load_all_data(self) -> Dict[str, List[Dict[str, Any]]]:
        """Read all records from Excel sheets into memory"""
        with self._lock:
            if not os.path.exists(self.file_path):
                return {k: [] for k in SHEET_SCHEMAS}

            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            result = {}

            for sheet_name, headers in SHEET_SCHEMAS.items():
                rows = []
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    # Read rows after header
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if not any(row):
                            continue
                        row_dict = {}
                        for idx, header in enumerate(headers):
                            val = row[idx] if idx < len(row) else None
                            row_dict[header] = val
                        rows.append(row_dict)
                result[sheet_name] = rows

            wb.close()
            return result

    def append_referral(self, ref: Dict[str, Any]) -> None:
        """Append a new referral row to Referrals sheet"""
        with self._lock:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb["Referrals"]

            now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
            row_values = [
                ref.get("referral_number", ""),
                ref.get("candidate_name", ""),
                ref.get("candidate_email", ""),
                ref.get("candidate_phone", ""),
                ref.get("referred_by_name", ""),
                ref.get("years_of_experience", 0.0),
                ref.get("relationship", ""),
                ref.get("position_title", ""),
                ref.get("status", "Submitted"),
                ref.get("id", ""),
                ref.get("position_id", ""),
                ref.get("referred_by_user_id", ""),
                ref.get("linkedin_url", "") or "",
                ref.get("github_url", "") or "",
                ref.get("original_filename", ""),
                ref.get("sharepoint_file_url", "") or "",
                ref.get("referral_note", ""),
                ref.get("created_at", now_str),
                ref.get("updated_at", now_str),
            ]
            ws.append(row_values)

            # Style appended row
            row_idx = ws.max_row
            for col_idx in range(1, len(row_values) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = BORDER_THIN
                cell.alignment = Alignment(vertical="center")

            # Also log to StatusHistory sheet
            ws_hist = wb["StatusHistory"]
            import uuid
            hist_row = [
                str(uuid.uuid4()),
                ref.get("id", ""),
                ref.get("referral_number", ""),
                "",
                ref.get("status", "Submitted"),
                ref.get("referred_by_name", "Employee"),
                f"Referral submitted by {ref.get('referred_by_name', 'Employee')}.",
                now_str,
            ]
            ws_hist.append(hist_row)

            self._auto_adjust_column_widths(ws)
            wb.save(self.file_path)
            wb.close()

    def update_referral_status(
        self,
        referral_id: str,
        referral_number: str,
        new_status: str,
        comment: Optional[str],
        changed_by: str,
    ) -> None:
        """Update referral status and append audit history"""
        with self._lock:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb["Referrals"]

            now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
            old_status = ""

            # Locate referral row by referral_id (Column 10: Referral ID)
            for row in range(2, ws.max_row + 1):
                cell_id = ws.cell(row=row, column=10).value
                if cell_id == referral_id:
                    old_status = ws.cell(row=row, column=9).value or ""
                    ws.cell(row=row, column=9).value = new_status  # Column 9: Status
                    ws.cell(row=row, column=19).value = now_str   # Column 19: Updated At
                    break

            # Append to StatusHistory
            import uuid
            ws_hist = wb["StatusHistory"]
            ws_hist.append([
                str(uuid.uuid4()),
                referral_id,
                referral_number,
                old_status,
                new_status,
                changed_by,
                comment or f"Status transitioned to {new_status}",
                now_str,
            ])

            # If status transitioned away from Hired, remove from HiredHistory worksheet
            if "HiredHistory" in wb.sheetnames and new_status != "Hired":
                ws_hh = wb["HiredHistory"]
                for r in range(2, ws_hh.max_row + 1):
                    if str(ws_hh.cell(row=r, column=11).value or "").strip() == str(referral_id).strip():
                        ws_hh.delete_rows(r)
                        break

            wb.save(self.file_path)
            wb.close()


    def append_hr_note(
        self,
        referral_id: str,
        referral_number: str,
        note: str,
        created_by: str,
    ) -> None:
        """Append a note to HRNotes sheet"""
        with self._lock:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb["HRNotes"]
            import uuid
            now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
            ws.append([
                str(uuid.uuid4()),
                referral_id,
                referral_number,
                created_by,
                note,
                now_str,
                now_str,
            ])
            wb.save(self.file_path)
            wb.close()

    def save_job_position(self, job: Dict[str, Any]) -> None:
        """Insert or update a job opening in JobPositions sheet"""
        with self._lock:
            wb = openpyxl.load_workbook(self.file_path)
            ws = wb["JobPositions"]

            job_id = job.get("id")
            found_row = None
            now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")

            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=1).value == job_id:
                    found_row = row
                    break

            row_data = [
                job_id,
                job.get("title", ""),
                job.get("department", ""),
                job.get("location", ""),
                job.get("employment_type", "Full-time"),
                "Yes" if job.get("is_active", True) else "No",
                job.get("description", ""),
                job.get("created_at", now_str),
                now_str,
            ]

            if found_row:
                for col_idx, val in enumerate(row_data, 1):
                    ws.cell(row=found_row, column=col_idx).value = val
            else:
                ws.append(row_data)

            self._auto_adjust_column_widths(ws)
            wb.save(self.file_path)
            wb.close()

    def get_workbook_bytes(self) -> bytes:
        """Return binary content of the Excel workbook"""
        with self._lock:
            if not os.path.exists(self.file_path):
                return b""
            with open(self.file_path, "rb") as fp:
                return fp.read()

    def delete_referral(self, referral_id: str) -> None:
        """Delete referral row from Referrals sheet and cleanup its history/notes in Excel"""
        with self._lock:
            if not os.path.exists(self.file_path):
                return
            wb = openpyxl.load_workbook(self.file_path)

            # 1. Delete from Referrals sheet
            if "Referrals" in wb.sheetnames:
                ws = wb["Referrals"]
                row_to_delete = None
                for r in range(2, ws.max_row + 1):
                    if str(ws.cell(row=r, column=10).value or "").strip() == str(referral_id).strip():
                        row_to_delete = r
                        break
                if row_to_delete:
                    ws.delete_rows(row_to_delete)

            # 2. Clean up StatusHistory sheet
            if "StatusHistory" in wb.sheetnames:
                ws_h = wb["StatusHistory"]
                rows_to_del = []
                for r in range(2, ws_h.max_row + 1):
                    if str(ws_h.cell(row=r, column=2).value or "").strip() == str(referral_id).strip():
                        rows_to_del.append(r)
                for r in reversed(rows_to_del):
                    ws_h.delete_rows(r)

            # 3. Clean up HRNotes sheet
            if "HRNotes" in wb.sheetnames:
                ws_n = wb["HRNotes"]
                rows_to_del = []
                for r in range(2, ws_n.max_row + 1):
                    if str(ws_n.cell(row=r, column=2).value or "").strip() == str(referral_id).strip():
                        rows_to_del.append(r)
                for r in reversed(rows_to_del):
                    ws_n.delete_rows(r)

            # 4. Clean up HiredHistory sheet
            if "HiredHistory" in wb.sheetnames:
                ws_hh = wb["HiredHistory"]
                row_to_del = None
                for r in range(2, ws_hh.max_row + 1):
                    if str(ws_hh.cell(row=r, column=11).value or "").strip() == str(referral_id).strip():
                        row_to_del = r
                        break
                if row_to_del:
                    ws_hh.delete_rows(row_to_del)

            wb.save(self.file_path)
            wb.close()
            logger.info(f"Deleted referral {referral_id} from Microsoft Excel workbook.")


    def save_user(self, user_data: Dict[str, Any]) -> None:
        """Insert or update user in Users worksheet"""
        with self._lock:
            if not os.path.exists(self.file_path):
                self.initialize_workbook()
            wb = openpyxl.load_workbook(self.file_path)
            if "Users" not in wb.sheetnames:
                ws = wb.create_sheet(title="Users")
                self._style_header_row(ws, SHEET_SCHEMAS["Users"])
            else:
                ws = wb["Users"]

            headers = SHEET_SCHEMAS["Users"]
            user_id = str(user_data.get("id") or "").strip()
            email = str(user_data.get("email") or "").strip().lower()

            target_row = None
            email_col = headers.index("Email") + 1
            id_col = headers.index("User ID") + 1

            for r in range(2, ws.max_row + 1):
                c_email = str(ws.cell(row=r, column=email_col).value or "").strip().lower()
                c_id = str(ws.cell(row=r, column=id_col).value or "").strip()
                if (email and c_email == email) or (user_id and c_id == user_id):
                    target_row = r
                    break

            row_num = target_row if target_row else ws.max_row + 1
            now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")

            row_values = [
                user_id,
                user_data.get("name", ""),
                user_data.get("email", ""),
                user_data.get("password", ""),
                user_data.get("role", "employee"),
                user_data.get("department", ""),
                user_data.get("created_at") or now_str,
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = BORDER_THIN

            self._auto_adjust_column_widths(ws)
            wb.save(self.file_path)
            wb.close()
            logger.info(f"Saved user {email} to Users worksheet (row {row_num}).")

    def get_user_by_email(self, email: str) -> Optional[Dict[str, Any]]:
        """Look up user directly from the Users worksheet in Excel"""
        with self._lock:
            if not os.path.exists(self.file_path):
                return None
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if "Users" not in wb.sheetnames:
                wb.close()
                return None
            ws = wb["Users"]
            sheet_headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
            email_idx = None
            for idx, h in enumerate(sheet_headers):
                if h.lower() == "email":
                    email_idx = idx
                    break

            if email_idx is None:
                wb.close()
                return None

            clean_email = email.strip().lower()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                row_email = str(row[email_idx] or "").strip().lower() if email_idx < len(row) else ""
                if row_email == clean_email:
                    res = {}
                    for idx, h in enumerate(sheet_headers):
                        if h:
                            res[h] = row[idx] if idx < len(row) else None
                    wb.close()
                    return res

            wb.close()
            return None

    def save_hired_record(self, hired: Dict[str, Any]) -> None:
        """Insert or update a candidate record in HiredHistory sheet"""
        with self._lock:
            if not os.path.exists(self.file_path):
                self.initialize_workbook()
            wb = openpyxl.load_workbook(self.file_path)
            if "HiredHistory" not in wb.sheetnames:
                ws = wb.create_sheet(title="HiredHistory")
                self._style_header_row(ws, SHEET_SCHEMAS["HiredHistory"])
            else:
                ws = wb["HiredHistory"]

            ref_id = str(hired.get("id") or hired.get("referral_id") or "").strip()
            existing_row = None
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=11).value or "").strip() == ref_id:
                    existing_row = r
                    break

            hired_at = hired.get("hired_at")
            if isinstance(hired_at, datetime):
                hired_at_str = hired_at.strftime("%Y-%m-%d %H:%M:%S")
            else:
                hired_at_str = str(hired_at or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"))

            row_values = [
                hired.get("referral_number", ""),
                hired.get("candidate_name", ""),
                hired.get("candidate_email", ""),
                hired.get("position_title", ""),
                hired.get("department", ""),
                hired.get("location", ""),
                hired.get("employment_type", "Full-time"),
                hired.get("referred_by_name", "Employee"),
                hired_at_str,
                hired.get("status", "Hired"),
                ref_id,
                str(hired.get("position_id") or ""),
            ]

            if existing_row:
                for col_idx, val in enumerate(row_values, start=1):
                    ws.cell(row=existing_row, column=col_idx, value=val)
            else:
                ws.append(row_values)
                r_idx = ws.max_row
                for col_idx in range(1, len(row_values) + 1):
                    cell = ws.cell(row=r_idx, column=col_idx)
                    cell.font = DATA_FONT
                    cell.border = BORDER_THIN
                    cell.alignment = Alignment(vertical="center")

            self._auto_adjust_column_widths(ws)
            wb.save(self.file_path)
            wb.close()
            logger.info(f"Saved hired candidate {hired.get('candidate_name')} ({ref_id}) to HiredHistory worksheet.")


def generate_hired_history_workbook_bytes(hired_items: List[Any]) -> bytes:
    """
    Generate a standalone styled Microsoft Excel (.xlsx) workbook for Hired Referral History.
    Can accept Referral models, dicts, or Pydantic schemas.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HiredHistory"

    headers = SHEET_SCHEMAS["HiredHistory"]
    ws.append(headers)
    ws.freeze_panes = "A2"
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN
    ws.row_dimensions[1].height = 26

    for item in hired_items:
        if isinstance(item, dict):
            ref_num = item.get("referral_number", "")
            cand_name = item.get("candidate_name", "")
            cand_email = item.get("candidate_email", "")
            pos_title = item.get("position_title", "")
            dept = item.get("department", "")
            loc = item.get("location", "")
            emp_type = item.get("employment_type", "Full-time")
            ref_by = item.get("referred_by_name", "")
            hired_at = item.get("hired_at", "")
            status_val = item.get("status", "Hired")
            ref_id = item.get("id") or item.get("referral_id", "")
            pos_id = item.get("position_id", "")
        else:
            ref_num = getattr(item, "referral_number", "")
            cand_name = getattr(item, "candidate_name", "")
            cand_email = getattr(item, "candidate_email", "")
            pos = getattr(item, "position", None)
            pos_title = getattr(item, "position_title", (pos.title if pos else ""))
            dept = getattr(item, "department", (pos.department if pos else ""))
            loc = getattr(item, "location", (pos.location if pos else ""))
            emp_type = getattr(item, "employment_type", (pos.employment_type if pos else "Full-time"))
            ref_by = getattr(item, "referred_by_name", "")
            hired_at = getattr(item, "hired_at", (getattr(item, "updated_at", None) or getattr(item, "created_at", None)))
            status_val = "Hired"
            ref_id = getattr(item, "id", "")
            pos_id = getattr(item, "position_id", "")

        if isinstance(hired_at, datetime):
            hired_at_str = hired_at.strftime("%Y-%m-%d %H:%M:%S")
        else:
            hired_at_str = str(hired_at or "")

        row = [
            ref_num,
            cand_name,
            cand_email,
            pos_title,
            dept,
            loc,
            emp_type or "Full-time",
            ref_by or "Employee",
            hired_at_str,
            status_val,
            str(ref_id),
            str(pos_id),
        ]
        ws.append(row)
        r_idx = ws.max_row
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()

